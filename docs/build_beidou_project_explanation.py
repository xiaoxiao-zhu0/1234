from __future__ import annotations

import json
import math
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
BUILD = DOCS / "_build_beidou_explanation"
OUTPUT = DOCS / "beidou_project_explanation.docx"
RESULT_PATH = ROOT / "data_contract" / "public_result.json"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
INK = "20262E"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF4D6"
GOLD = "9A6A00"
PALE_RED = "FDECEC"
RED = "9B1C1C"
GREEN = "26734D"
WHITE = "FFFFFF"

BODY_FONT = "Microsoft YaHei"
TITLE_FONT = "Microsoft YaHei"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=None, bold=None, color=INK, italic=None, font=BODY_FONT) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_paragraph_keep(p, keep_next=False, keep_lines=True) -> None:
    p_pr = p._p.get_or_add_pPr()
    if keep_next:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    if keep_lines:
        node = OxmlElement("w:keepLines")
        p_pr.append(node)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def create_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1

    def make(fmt: str, text_value: str, abstract_id: int, num_id: int) -> int:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text_value)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), BODY_FONT)
        r_fonts.set(qn("w:hAnsi"), BODY_FONT)
        r_fonts.set(qn("w:eastAsia"), BODY_FONT)
        r_pr.append(r_fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(abstract_id))
        num.append(abs_ref)
        numbering.append(num)
        return num_id

    bullet_id = make("bullet", "•", next_abs, next_num)
    decimal_id = make("decimal", "%1.", next_abs + 1, next_num + 1)
    return bullet_id, decimal_id


def apply_num(p, num_id: int) -> None:
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


def add_body(doc, text: str, bold_lead: str | None = None, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, bold=True, color=NAVY)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_bullet(doc, text: str, bullet_id: int, bold_lead: str | None = None):
    p = doc.add_paragraph()
    apply_num(p, bullet_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, size=10.5, bold=True, color=NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_number(doc, text: str, number_id: int, bold_lead: str | None = None):
    p = doc.add_paragraph()
    apply_num(p, number_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, size=10.5, bold=True, color=NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_callout(doc, label: str, text: str, fill=LIGHT_BLUE, color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, bold=True, color=color)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, font=TITLE_FONT)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], header_fill=LIGHT_GRAY,
              aligns: list[str] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.rows[0].cells[0]
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, bold=True, color=NAVY)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            align = (aligns[idx] if aligns else "left")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=9.1)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def set_style(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = styles[style_name]
        style.font.name = TITLE_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), TITLE_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), TITLE_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), TITLE_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("北斗巡知 | 项目完整说明书")
    set_run_font(r, size=8.5, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("内部沟通与竞赛讲解版  |  ")
    set_run_font(r, size=8.5, color=MUTED)
    add_field(p, "PAGE")


def find_font(size: int, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    lines, current = [], ""
    for ch in text:
        trial = current + ch
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_box(draw, xy, title, body, fill, outline=BLUE, title_color=NAVY):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=3)
    title_font = find_font(30, True)
    body_font = find_font(23)
    draw.text((x1 + 24, y1 + 18), title, font=title_font, fill=title_color)
    y = y1 + 70
    for line in wrap(draw, body, body_font, x2 - x1 - 48):
        draw.text((x1 + 24, y), line, font=body_font, fill=INK)
        y += 36


def create_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1500, 900), "white")
    d = ImageDraw.Draw(img)
    title = find_font(38, True)
    d.text((70, 35), "北斗时空信息与视觉识别的双路协同架构", font=title, fill=NAVY)
    draw_box(d, (70, 130, 600, 300), "视觉数据流", "航拍图像 → 目标识别模型 → 类别、置信度、语义特征", LIGHT_BLUE)
    draw_box(d, (70, 390, 600, 590), "北斗时空数据流", "经纬度、时间、定位质量 → 时空索引、区域划分 → location_id", "EEF7EF", GREEN)
    draw_box(d, (850, 205, 1430, 465), "关联与诊断", "sample_id + timestamp 关联两路数据；计算分区域识别率、遗忘率、定位误差和综合风险", LIGHT_GRAY, DARK_BLUE)
    draw_box(d, (850, 570, 1430, 800), "业务决策输出", "回放预算建议、重点复飞坐标、补采清单、人工复核任务、诊断报告", PALE_GOLD, GOLD)
    arrow_font = find_font(48, True)
    d.text((675, 255), "→", font=arrow_font, fill=BLUE)
    d.text((675, 470), "→", font=arrow_font, fill=GREEN)
    d.text((1095, 490), "↓", font=arrow_font, fill=GOLD)
    small = find_font(22)
    d.text((70, 830), "核心原则：图像回答“看到了什么”；北斗回答“在哪里、什么时候、哪个区域持续识别不好，以及应该去哪里重新巡检”。", font=small, fill=MUTED)
    img.save(path, quality=95)


def create_loops_diagram(path: Path) -> None:
    img = Image.new("RGB", (1500, 940), "white")
    d = ImageDraw.Draw(img)
    title = find_font(38, True)
    d.text((70, 35), "产品运行由两个闭环组成", font=title, fill=NAVY)
    draw_box(d, (60, 120, 720, 830), "日常巡检闭环", "1. 创建巡检任务\n2. 无人机按航线采集图像和北斗元数据\n3. 后端校验并按地理区域归档\n4. 当前模型执行识别\n5. 计算区域风险和定位质量\n6. 生成复飞、补采和人工复核任务\n7. 导出巡检诊断报告", LIGHT_BLUE)
    draw_box(d, (780, 120, 1440, 830), "模型更新闭环", "1. 收集低置信度、未知目标和误识别样本\n2. 人工确认是否形成新类别\n3. 构建新的类别增量任务\n4. 私有三层持续学习引擎在服务器2训练\n5. 与当前模型、ER基线比较\n6. 检查准确率、遗忘率和高风险区域\n7. 仅发布通过门槛的模型版本", "EEF7EF", GREEN)
    note = find_font(22, True)
    d.text((145, 865), "日常上传不会自动触发训练", font=note, fill=RED)
    d.text((930, 865), "模型更新是受控、可审核、可回滚的版本流程", font=note, fill=GREEN)
    img.save(path, quality=95)


def create_isolation_diagram(path: Path) -> None:
    img = Image.new("RGB", (1500, 850), "white")
    d = ImageDraw.Draw(img)
    title = find_font(38, True)
    d.text((70, 35), "团队协作与核心代码隔离", font=title, fill=NAVY)
    draw_box(d, (70, 150, 620, 680), "算法负责人私有区", "服务器2 / GPU1\n• 私有三层算法源码\n• 训练配置与模型权重\n• 原始预测文件与日志\n• 统一评测和质量审核\n\n对外只导出白名单结果 JSON", PALE_RED, RED)
    draw_box(d, (880, 150, 1430, 680), "共享竞赛工程", "GitHub 私有仓库\n• 前端界面与交互\n• 业务仿真和产品图\n• public_result.json\n• 文档、BP、PPT、测试\n\n队员及其 AI 可完整优化共享区", LIGHT_BLUE, BLUE)
    arrow_font = find_font(44, True)
    d.text((650, 350), "审核后 JSON  →", font=arrow_font, fill=GOLD)
    note = find_font(23)
    d.text((180, 735), "共享工程能让队员理解完整产品流程，但不能反推出私有网络、损失函数、回放采样或模型权重。", font=note, fill=MUTED)
    img.save(path, quality=95)


def add_picture(doc, path: Path, caption: str, width=6.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_paragraph_keep(p, keep_next=True)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(0)
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption)
    set_run_font(r, size=8.8, color=MUTED)
    return p


def load_public_result():
    with RESULT_PATH.open("r", encoding="utf-8") as f:
        return json.load(f)


def build_document():
    BUILD.mkdir(parents=True, exist_ok=True)
    architecture = BUILD / "architecture.png"
    loops = BUILD / "loops.png"
    isolation = BUILD / "isolation.png"
    create_architecture_diagram(architecture)
    create_loops_diagram(loops)
    create_isolation_diagram(isolation)

    result = load_public_result()
    methods = {m["name"]: m for m in result["methods"]}
    er = next(m for m in result["methods"] if "ER" in m["name"])
    primary = next(m for m in result["methods"] if m["name"] == result["primary_method"])
    locations = result["locations"]

    doc = Document()
    set_style(doc)
    bullet_id, number_id = create_numbering(doc)

    # Cover: editorial_cover pattern with restrained technical-report treatment.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(94)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第十五届中国创新创业大赛\n北斗时空信息专业赛")
    set_run_font(r, size=13, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("北斗巡知")
    set_run_font(r, size=30, bold=True, color=NAVY, font=TITLE_FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("北斗时空感知持续学习技术与低空无人机巡检系统")
    set_run_font(r, size=16, color=DARK_BLUE, font=TITLE_FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(70)
    r = p.add_run("项目完整说明书｜团队统一认知与评审讲解版")
    set_run_font(r, size=11.5, color=MUTED)

    add_callout(doc, "项目一句话", "面向低空无人机常态化巡检，以北斗时空索引连接视觉识别、局部风险诊断和复飞决策，并用类别增量持续学习应对巡检目标持续新增和历史知识遗忘。", fill=LIGHT_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(45)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("版本：V1.0  |  日期：2026年8月5日")
    set_run_font(r, size=10, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("证据状态：公开航拍数据实验 + 虚拟北斗坐标仿真；非真实外业验收")
    set_run_font(r, size=9.5, bold=True, color=RED)
    doc.add_page_break()

    add_heading(doc, "阅读说明", 1)
    add_body(doc, "本说明书用于统一团队对项目的理解，并为商业计划书、项目申报、产品界面、路演PPT和评委答辩提供一致口径。文档以当前共享竞赛工程和已审核公开结果为依据，重点解释项目到底解决什么问题、北斗为什么不可替代、算法如何转化为业务能力，以及当前原型与完整产品之间还差哪些环节。")
    add_callout(doc, "最重要的边界", "本文不会公开或推测私有三层持续学习算法的网络结构、损失函数、回放采样、训练配置和模型权重。对外只描述已经验证的能力、输入输出关系和业务价值。", fill=PALE_RED, color=RED)

    add_heading(doc, "阅读导航", 2)
    add_table(doc, ["阅读对象", "建议重点阅读"], [
        ["全体队员", "执行摘要、项目定位、完整业务闭环、当前实现状态"],
        ["前端与产品成员", "用户角色、任务数据包、前后端职责、页面与真实产品的差距"],
        ["算法负责人", "双路数据架构、模型更新闭环、实验结果解释、证据边界"],
        ["BP与PPT成员", "创新性、商业价值、评委口径、禁止表述与FAQ"],
    ], [2200, 7160])

    add_heading(doc, "执行摘要", 1)
    add_body(doc, "“北斗巡知”不是一个把持续学习算法套进无人机外壳的视觉分类演示，而是一套面向常态化低空巡检的时空风险诊断与持续学习平台。它将无人机采集的图像与北斗位置、时间和定位质量通过样本编号关联，先回答“看到了什么”，再回答“在哪里识别不好、哪些区域正在遗忘、有限模型维护资源应该投向哪里、下一次应该飞到哪里复核”。")
    add_body(doc, "项目当前最适合申报“北斗时空信息应用—产业与经济应用—低空经济（无人机巡检）”，并可关联公共基础设施中的能源、电网和通信基站应用。它不应申报为北斗定位导航核心技术或芯片终端设备，因为当前创新重点不在定位算法、导航芯片和硬件国产化，而在北斗时空信息与人工智能持续学习、风险治理和巡检作业闭环的融合。")
    add_body(doc, f"当前公开数据实验采用VisDrone目标裁剪分类数据、10个类别、5个顺序任务、4301条元数据和预测记录、固定600条回放预算。代表性随机种子下，ER经验回放最终准确率为{er['average_accuracy']*100:.2f}%，三层方法为{primary['average_accuracy']*100:.2f}%，提升{(primary['average_accuracy']-er['average_accuracy'])*100:.2f}个百分点；平均遗忘率由{er['average_forgetting']*100:.2f}%降至{primary['average_forgetting']*100:.2f}%，下降约{(er['average_forgetting']-primary['average_forgetting'])*100:.2f}个百分点。该结果证明当前方法在既定公开数据、任务划分和预算约束下，比ER更能保持历史类别知识；它不等于真实电网外业准确率，也不证明已经完成机载部署。")
    add_callout(doc, "北斗不可替代的作用", "图像回答“看到了什么”；北斗回答“在哪里、什么时候、哪个区域持续识别不好，以及应该去哪里重新巡检”。", fill=PALE_GOLD, color=GOLD)

    add_heading(doc, "1. 项目定位与参赛方向", 1)
    add_heading(doc, "1.1 推荐定位", 2)
    add_callout(doc, "正式定位", "北斗时空信息应用方向下，面向低空经济和能源基础设施巡检的时空风险诊断与持续学习平台。")
    add_body(doc, "项目的主应用场景是搭载北斗定位模块的无人机，对输电走廊、山区边坡、施工扰动区、通信塔区、光伏场站等区域进行长期巡检。随着巡检时间拉长，系统会不断遇到新的车辆类型、施工机械、人员行为、围栏缺陷或其他异常目标，因此模型不能依赖每次收集全部历史数据后重新训练，而需要在有限存储和计算预算下持续吸收新类别，同时尽量保持旧类别能力。")

    add_heading(doc, "1.2 对应赛事分类", 2)
    add_table(doc, ["层级", "项目归属", "判断理由"], [
        ["一级方向", "二、北斗时空信息应用", "核心是将北斗时空信息用于巡检业务和模型治理，而非研制北斗芯片或定位算法。"],
        ["主申报条目", "（二）产业与经济应用—低空经济—无人机巡检", "产品围绕无人机巡航、风险定位、重点复飞、补采和模型持续更新形成闭环。"],
        ["关联条目", "（一）公共基础设施应用—能源、电网、通信基站", "首个业务叙事选择输电走廊和基础设施通道巡检。"],
        ["技术融合标签", "北斗时空信息 + 人工智能 + 大数据/云计算", "北斗提供时空索引，持续学习提供模型更新，后端沉淀跨架次区域风险。"],
    ], [1500, 3100, 4760])

    add_heading(doc, "1.3 为什么不主报“关键技术”", 2)
    add_body(doc, "赛事中的北斗关键技术通常要求在定位、导航、授时、通导遥融合、短报文增强、高精度或高可靠安全方面形成直接技术突破。当前项目没有提出新的北斗定位解算方法、增强定位链路、授时算法、芯片模组或天线，因此若主报关键技术，评委很容易追问定位精度提升、硬件指标和实测认证，项目证据无法支撑。更准确的说法是：项目属于北斗应用创新，人工智能算法是关键支撑技术。")

    add_heading(doc, "2. 为什么不是“增量学习算法加无人机外壳”", 1)
    add_body(doc, "一个普通的类别增量学习算法，只回答模型如何连续学习新类别；一个真正适配北斗巡检的系统，还必须回答新数据来自哪次任务、发生在哪个坐标、哪个区域风险最高、定位质量是否异常、有限回放预算如何配置，以及识别失效后如何触发复飞和人工复核。项目的场景真实性不来自页面上放置无人机图标，而来自时空数据进入了诊断和决策闭环。")
    add_table(doc, ["判断维度", "普通持续学习演示", "北斗巡知"], [
        ["输入", "图像与类别标签", "图像 + 时间 + 北斗坐标 + 定位质量 + 航线/任务信息"],
        ["评价方式", "只看全局平均准确率", "同时看分区域识别率、遗忘率、定位误差和业务后果"],
        ["资源分配", "全局均匀回放或固定采样", "在固定总预算下给出区域风险预算建议"],
        ["业务输出", "得到一个新模型", "得到模型版本、薄弱区域、复飞坐标、补采与人工复核任务"],
        ["北斗价值", "可有可无的标签", "跨架次时空索引、区域诊断和现场任务定位的基础"],
    ], [1800, 3300, 4260])
    add_callout(doc, "判定标准", "去掉北斗后，系统仍能做图像分类，但无法可靠完成跨架次位置关联、区域性能诊断和坐标级复飞任务。因此北斗不是装饰，而是业务闭环的基础设施。")

    add_heading(doc, "3. 目标业务场景", 1)
    add_heading(doc, "3.1 首选场景：输电走廊外部环境巡检", 2)
    add_body(doc, "首选叙事应保持聚焦：运维单位定期使用搭载北斗模块的无人机巡检输电走廊及周边区域，识别人员、汽车、卡车、施工车辆、堆物、围栏和其他影响通道安全的目标。系统记录目标发生的时间与位置，持续评估各地理区域的识别稳定性，并把高风险坐标转化为复飞、补采和人工复核任务。")
    add_body(doc, "这个场景同时满足三项条件：一是存在明确的北斗位置需求；二是巡检是长期重复而不是一次性识别，天然形成持续学习任务流；三是不同区域的业务后果不同，同样的模型错误发生在输电走廊、山区边坡和普通农田缓冲区，其处置优先级不应相同。")

    add_heading(doc, "3.2 可扩展场景", 2)
    for text in [
        "交通与铁路沿线巡检：识别施工车辆、人员入侵、边坡风险和异物，并记录可复核坐标。",
        "通信基站与管线巡检：定位设备周边异常活动和环境变化，形成跨周期风险档案。",
        "自然资源与园区巡检：持续吸收新增地物和异常类别，减少模型因环境变化产生的局部失效。",
        "应急巡查：在灾后或突发事件中快速建立新增目标类别，并根据区域风险安排复飞与人工确认。",
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "4. 用户、客户与使用边界", 1)
    add_table(doc, ["角色", "主要任务", "使用界面/输出"], [
        ["无人机巡检操作员", "创建任务、执行航线、上传任务包、接收复飞坐标", "任务创建、上传进度、航线地图、复飞清单"],
        ["电网/设施运维管理人员", "查看区域风险、安排人工复核、导出报告", "风险地图、区域诊断、告警与报告"],
        ["算法与模型运维人员", "确认新类别、发起受控训练、评估候选模型、发布版本", "样本审核、模型对比、质量门槛、版本管理"],
        ["项目管理与决策人员", "评估巡检覆盖、模型稳定性和资源投入", "汇总看板、趋势指标、重点区域与处置闭环"],
    ], [1900, 3600, 3860])
    add_body(doc, "当前网页主要承担竞赛原型展示和结果解释，还不是完整交付给一线操作员的生产系统。完整产品应当让上述不同角色通过权限分层使用同一后端，而不是让所有用户直接操作训练代码。")

    add_heading(doc, "5. 场景原生痛点", 1)
    pain_points = [
        ("持续出现新类别", "巡检对象和风险类型随时间变化，模型会依次遇到未见类别。每次汇总全部历史数据重新训练成本高、周期长，也不适合频繁更新。"),
        ("不同位置分布不一致", "山区、河谷、城区、光伏场站在光照、视角、背景和目标密度上差异明显，全局平均准确率可能掩盖某些区域已经失效。"),
        ("定位质量与跨架次关联", "北斗定位误差、遮挡和多路径效应会影响样本归属区域，也会影响后续能否回到同一地点复核。"),
        ("经验回放存在局部失效", "均匀保存历史样本不代表每个高风险区域都获得足够训练曝光，稀少但重要的区域更容易被遗忘。"),
        ("容易学习虚假关联", "视角、光照、地貌和位置共现关系可能被模型当成类别依据，导致跨区域泛化下降。"),
        ("存储和训练预算有限", "机载端与服务器维护都需要控制样本、特征、算力和更新时间，不能无限保存原图或无限扩充回放缓存。"),
    ]
    for idx, (title, body) in enumerate(pain_points, 1):
        add_heading(doc, f"5.{idx} {title}", 2)
        add_body(doc, body)

    add_heading(doc, "6. 数据从哪里来", 1)
    add_heading(doc, "6.1 商业运行中的真实数据来源", 2)
    add_body(doc, "未来商业系统的数据主要来自客户自己的无人机巡检任务，而不是让用户手工向分类模型输入经纬度。一次任务结束后，无人机或任务平台上传一个标准任务包，后端通过sample_id和timestamp把图像与北斗元数据关联。")
    add_table(doc, ["字段", "含义", "是否进入视觉分类网络"], [
        ["mission_id / route_id", "巡检任务和航线标识", "否，用于任务管理与跨架次追踪"],
        ["image / video", "相机采集的图像或视频帧", "是，属于视觉识别输入"],
        ["sample_id", "样本唯一编号", "否，用于数据关联和审计"],
        ["timestamp", "采集时间", "否，用于时序关联与版本分析"],
        ["latitude / longitude", "北斗经纬度", "暂不直接输入分类器，用于时空索引和区域诊断"],
        ["altitude", "飞行高度", "可作为诊断元数据，当前不作为核心分类输入"],
        ["position_quality", "定位质量或误差估计", "否，用于过滤、风险解释和复飞可信度"],
        ["human_label", "可选人工确认标签", "用于训练与评测，但需要审核"],
    ], [2200, 4400, 2760])

    add_heading(doc, "6.2 当前比赛仿真数据来源", 2)
    add_body(doc, "当前阶段使用公开VisDrone航拍目标检测数据，将目标框裁剪为分类样本，按新类别出现顺序构建5个类别增量任务；再为样本生成虚拟巡检顺序、WGS84坐标、8个地理区域和高斯定位扰动，从而低成本复现“无人机沿航线持续采集—不同区域表现不同—新类别陆续出现”的业务结构。")
    add_callout(doc, "合规表述", "公开数据是真实图像数据；经纬度、航线、定位噪声和业务区域是仿真构造。可以称为“北斗无人机巡检仿真验证”，不能称为“真实北斗外业实测”。", fill=PALE_RED, color=RED)

    add_heading(doc, "7. 完整端到端流程", 1)
    add_picture(doc, architecture, "图1  北斗时空信息与视觉识别采用双路协同，而非把经纬度强行输入分类器")
    add_body(doc, "完整系统不是一个单页面，也不是一个训练脚本，而是从任务创建、数据采集、识别诊断到模型更新和复飞处置的完整链路。前端是用户入口，业务后端负责数据与任务编排，算法后端负责推理、持续学习和质量评估。")
    workflow = [
        "运维人员在平台创建巡检任务，设置区域、航线、目标类型和任务时间。",
        "搭载北斗模块的无人机沿预定航线飞行，采集图像、时间、坐标、高度和定位质量。",
        "任务包上传后，业务后端校验字段完整性，并依据sample_id和timestamp完成图像与时空元数据关联。",
        "时空索引模块将样本分配到地理网格或业务区域，保留跨架次可追踪的位置身份。",
        "当前模型对图像执行识别，输出类别、置信度和必要的公开诊断字段。",
        "诊断模块按区域计算识别率、遗忘率、定位误差、样本量和模型不稳定性。",
        "风险模块把模型不稳定性与区域业务后果结合，在固定总预算下给出优先级和资源建议。",
        "系统生成复飞坐标、补采清单和人工复核任务，并在前端形成可解释报告。",
        "经过人工确认的新类别进入受控模型更新闭环，候选模型通过门槛后再发布。",
    ]
    for item in workflow:
        add_number(doc, item, number_id)

    add_heading(doc, "8. 两个闭环必须分开理解", 1)
    add_picture(doc, loops, "图2  日常巡检闭环负责业务运行，模型更新闭环负责受控学习")
    add_heading(doc, "8.1 日常巡检闭环", 2)
    add_body(doc, "日常巡检的目标是及时完成识别、定位和处置。用户上传一次任务数据后，系统首先执行数据校验、推理和区域诊断，而不是立即启动模型训练。这样可以避免未经确认的新样本污染模型，也便于保持生产模型版本稳定。")
    add_heading(doc, "8.2 模型更新闭环", 2)
    add_body(doc, "模型更新由算法或模型运维人员发起。系统先汇总低置信度、未知目标和误识别样本，再由人工确认是否形成新类别；随后在服务器2的GPU1上运行私有三层持续学习引擎。候选模型必须与当前模型和ER基线在统一任务划分、统一预算和统一评测口径下比较，只有通过准确率、遗忘率和高风险区域表现门槛后才允许发布。")
    add_callout(doc, "产品原则", "训练不是每次上传后的自动副作用，而是一套可审核、可比较、可回滚的模型版本流程。")

    add_heading(doc, "9. 北斗与视觉模型的正确关系", 1)
    add_heading(doc, "9.1 当前为什么不直接把经纬度输入分类网络", 2)
    add_body(doc, "经纬度与目标类别可能存在强烈的偶然共现。例如某段坐标在训练期经常出现卡车，模型可能把“位置”当成“卡车”的捷径；当卡车出现在另一区域，或原区域环境发生变化时，模型就会产生错误。直接输入坐标还可能让定位噪声传播到分类结果，使模型对轻微位置偏移过度敏感。")
    add_body(doc, "因此，当前更稳妥的工程方案是双路协同：图像进入视觉模型；北斗坐标、时间和定位质量进入时空诊断模块；两路通过sample_id和timestamp关联。北斗信息影响区域统计、风险排序、回放预算建议和复飞任务，但暂不作为目标类别预测的直接特征。")
    add_heading(doc, "9.2 什么时候可以研究坐标辅助输入", 2)
    add_body(doc, "未来可以在有真实跨区域数据、严格消融实验和位置泄漏防护的前提下，研究弱位置先验或地理上下文分支。但这属于后续研究，不应在当前BP中写成已完成能力。即便加入，也必须证明提升来自可迁移的时空上下文，而不是记住训练地点。")

    add_heading(doc, "10. 私有三层算法在产品中的位置", 1)
    add_body(doc, "共享竞赛工程中的“外部私有持续学习引擎”由算法负责人独立维护。团队对外可以使用“三层因果语义回放”作为已审核的方法名称，但不应在没有论文定稿和知识产权安排的情况下披露内部结构。产品说明只需要清楚表达它承担的三类能力：在类别持续增加时吸收新知识、通过因果语义机制降低干扰因素造成的虚假关联、在有限回放预算下减少历史知识遗忘。")
    add_table(doc, ["公开可讲能力", "产品意义", "当前证据"], [
        ["类别增量持续学习", "新类别出现时无需每次汇总全部历史数据从头训练", "10类、5个顺序任务的公开数据实验"],
        ["因果语义回放", "强化目标本体语义，降低光照、视角和背景共现带来的干扰", "以主方法整体准确率和遗忘率结果体现；内部实现不公开"],
        ["位置特定诊断与风险预算", "识别局部薄弱区域，并在固定总量下提出区域资源配置建议", "8个虚拟区域的预测级诊断和600预算守恒"],
    ], [2400, 4000, 2960])
    add_callout(doc, "避免错误描述", "不要把产品功能层直接等同于论文网络的“三层结构”，也不要由共享前端反推算法实现。评委需要理解的是能力闭环和结果证据，而不是看到源码。", fill=PALE_RED, color=RED)

    add_heading(doc, "11. 系统架构与前后端职责", 1)
    add_table(doc, ["层级", "负责内容", "当前状态"], [
        ["用户前端", "任务创建、数据上传、地图风险、区域详情、新类别确认、复飞任务、报告导出", "已完成仿真页、结果看板和正式报告页；任务管理等仍待实现"],
        ["业务后端", "任务包校验、图像坐标关联、任务队列、权限、模型版本、API/JSON输出", "当前以静态服务和公开JSON契约为主，尚非生产后端"],
        ["算法后端", "图像推理、私有持续学习、候选模型评估、区域诊断、风险预算计算", "私有训练侧与公开结果导出链路已分离；未开放给队员"],
        ["数据与存储", "任务元数据、时空索引、预测、模型版本、报告和审计记录", "当前公开仓库只保留白名单结果与仿真数据"],
    ], [1600, 4750, 3010])
    add_body(doc, "后端是计算和治理核心，前端是用户直接操作和理解系统的入口。竞赛阶段前端具有重要作用：它把抽象算法指标转化为地图、风险区域、预算分配和复飞任务，让评委看到算法如何形成业务价值。但不能因为前端做得完整，就把静态仿真误写成真实在线产品。")

    add_heading(doc, "12. 当前前端究竟做了什么", 1)
    add_table(doc, ["页面", "当前功能", "证据类型", "不能宣称"], [
        ["交互仿真页 frontend/index.html", "模拟航线、定位扰动、8个区域、风险、预算和复飞队列；支持交互演示", "simulation_only", "不能称为真实模型在线训练或真实北斗航迹"],
        ["结果看板 frontend/dashboard.html", "读取public_result.json，展示方法指标、区域结果和预算分配", "prediction_backed", "不能推导真实外业、机载部署或客户验收"],
        ["正式报告页 frontend/report.html", "基于审核JSON生成可打印的诊断报告", "沿用正式JSON状态", "不能加入JSON之外的真实指标"],
    ], [2200, 3500, 1500, 2160])

    screenshot = ROOT / "assets" / "screenshots" / "beidou_product_board.png"
    if screenshot.exists():
        add_picture(doc, screenshot, "图3  当前竞赛产品板：用于解释航线、区域风险、方法结果与证据状态", width=6.15)

    add_heading(doc, "12.1 完整产品还需要补齐的前端能力", 2)
    for text in [
        "任务中心：创建巡检任务、选择区域和航线、查看任务状态与历史架次。",
        "数据接入：上传图像/视频与CSV或标准任务包，显示字段校验和关联失败原因。",
        "地图诊断：在真实地图或GIS底图上查看区域风险、定位质量、样本覆盖和历史趋势。",
        "样本审核：查看低置信度和未知目标，完成人工确认、新类别命名与审核。",
        "模型版本：展示当前模型、候选模型、指标门槛、发布记录和回滚入口。",
        "处置闭环：生成复飞任务、补采清单、人工复核工单并跟踪完成状态。",
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "13. 当前仿真如何计算", 1)
    add_heading(doc, "13.1 仿真任务流", 2)
    steps = [
        "从VisDrone检测数据中裁剪目标实例，形成10类航拍目标分类样本。",
        "按照新类别陆续出现的顺序划分为5个任务，模拟无人机长期巡检中目标集合扩展。",
        "按虚拟蛇形航线给样本分配时间和经纬度，并加入定位扰动，模拟北斗位置偏差。",
        "根据坐标将样本划入8个业务区域，例如输电走廊、山区边坡和施工扰动区。",
        "ER和三层方法使用相同任务流、相同预算和相同评测数据运行，输出逐样本预测。",
        "诊断程序按location_id汇总识别率、历史正确状态、遗忘率、定位误差和样本量。",
        "综合模型不稳定性与业务后果形成区域风险，并在总预算600不变的条件下分配建议额度。",
    ]
    for step in steps:
        add_number(doc, step, number_id)

    add_heading(doc, "13.2 区域风险的业务解释", 2)
    add_body(doc, "当前公开结果把区域风险理解为“模型不稳定性 × 业务后果”的综合量。模型不稳定性来源于识别率下降、遗忘、定位质量和样本表现等诊断信息；业务后果表示该区域一旦漏检可能造成的影响。它不是通用安全概率，也不是事故发生概率，而是用于排序补采、复飞和模型维护优先级的内部决策指标。")
    add_callout(doc, "当前实现边界", "位置风险已经参与区域诊断和预算建议，但公开材料应谨慎区分“预算建议已计算”与“预算已反向控制私有训练缓存”。后者需要在私有训练侧完成联调和对照实验后才能宣称。", fill=PALE_RED, color=RED)

    add_heading(doc, "14. 实验设置与正式结果", 1)
    add_table(doc, ["项目", "当前正式设置"], [
        ["数据集", "VisDrone-DET目标裁剪分类 + 模拟北斗坐标"],
        ["数据规模", f"{result['dataset']['metadata_count']}条元数据、{result['dataset']['prediction_count']}条预测记录"],
        ["类别/任务", f"{result['dataset']['class_count']}类、{result['dataset']['task_count']}个顺序类别增量任务"],
        ["对比方法", "ER经验回放 vs. 三层因果语义回放"],
        ["回放预算", f"固定总预算{result['budget']['total']}，分配总和{result['budget']['allocated_sum']}"],
        ["地理区域", f"{len(locations)}个虚拟业务区域，含定位误差、定位质量、遗忘和风险指标"],
        ["结果状态", f"{result['quality']['status']} / recommended_for_bp = {str(result['quality']['recommended_for_bp']).lower()}"],
    ], [2200, 7160])

    add_heading(doc, "14.1 方法对比", 2)
    add_table(doc, ["方法", "最终平均准确率", "宏平均准确率", "平均遗忘率"], [
        [er["name"], f"{er['average_accuracy']*100:.2f}%", f"{er['macro_accuracy']*100:.2f}%", f"{er['average_forgetting']*100:.2f}%"],
        [primary["name"], f"{primary['average_accuracy']*100:.2f}%", f"{primary['macro_accuracy']*100:.2f}%", f"{primary['average_forgetting']*100:.2f}%"],
        ["差异", f"+{(primary['average_accuracy']-er['average_accuracy'])*100:.2f} pp", f"+{(primary['macro_accuracy']-er['macro_accuracy'])*100:.2f} pp", f"-{(er['average_forgetting']-primary['average_forgetting'])*100:.2f} pp"],
    ], [3400, 1980, 1980, 2000], aligns=["left", "center", "center", "center"])
    add_body(doc, "这里的“准确率”用于评估模型在连续学习全部任务后，能否同时识别已经学过的新旧类别；“遗忘率”用于衡量旧任务在后续学习后下降了多少。对于持续学习项目，只有当前任务准确率高并不够，必须证明模型没有把前面学过的类别大规模忘掉。")

    add_heading(doc, "14.2 区域风险与预算", 2)
    loc_rows = []
    for loc in locations:
        loc_rows.append([
            f"{loc['location_id']} {loc['name']}",
            str(loc["sample_count"]),
            f"{loc['accuracy']*100:.2f}%",
            f"{loc['forgetting']*100:.2f}%",
            f"{loc['risk']:.3f}",
            str(loc["replay_budget"]),
        ])
    add_table(doc, ["区域", "样本数", "识别率", "遗忘率", "风险", "预算建议"], loc_rows,
              [2600, 1100, 1400, 1400, 1300, 1560], aligns=["left", "center", "center", "center", "center", "center"])
    top3 = sum(loc["replay_budget"] for loc in locations[:3])
    add_body(doc, f"风险最高的三个区域是G02山区边坡、G01输电走廊和G07施工扰动区，合计获得{top3}/{result['budget']['total']}={top3/result['budget']['total']*100:.0f}%的预算建议。这个结果用于说明固定总量下的资源倾斜逻辑，不代表真实客户已经按该额度完成训练或复飞。")

    add_heading(doc, "15. 43.59%准确率对项目有什么作用", 1)
    add_body(doc, "43.59%不是一个可以脱离任务难度单独判断“高或低”的产品指标。它的价值首先来自可比性：在同一公开数据、同一10类5任务、同一固定预算和同一评测流程下，三层方法比ER高12.11个百分点，同时遗忘率低19.23个百分点。这支持“方法在持续学习场景下比普通经验回放更能保留旧知识”的结论。")
    add_body(doc, "它还不能支持三类结论：不能说真实电网巡检准确率达到43.59%；不能说已经满足商业部署阈值；不能用一个随机种子代表所有运行都稳定。更强的BP证据应增加多随机种子均值和标准差、类别级结果、任务级曲线、消融实验，以及位置风险预算真正接入训练后的对照结果。")
    add_table(doc, ["可以支持", "暂时不能支持"], [
        ["三层方法在当前统一设置下优于ER", "真实外业识别精度或漏检率"],
        ["持续学习中的灾难性遗忘得到明显缓解", "机载实时性能、功耗和存储指标"],
        ["可以按区域发现局部性能差异并提出预算建议", "预算闭环已经在生产训练中自动生效"],
        ["公开数据仿真链路已经形成可展示原型", "客户验收、商业收入或规模化部署"],
    ], [4680, 4680])

    add_heading(doc, "16. 项目的创新性", 1)
    add_heading(doc, "16.1 时空感知的局部风险诊断", 2)
    add_body(doc, "传统持续学习评测通常只看全局平均准确率，而巡检业务更关心具体地点是否正在失效。项目利用北斗坐标建立location_id，在区域粒度统计识别率、遗忘率、定位误差和模型不稳定性，把“模型整体还可以”进一步拆解为“哪些地点已经不可靠”。")
    add_heading(doc, "16.2 风险预算而非均匀回放", 2)
    add_body(doc, "在固定回放总量下，根据区域风险和业务后果给出资源配置建议，高风险区域获得更多历史知识留存或补采资源。创新点不只是改变采样比例，而是把模型维护预算与可解释的地理风险连接起来。当前已完成预算建议计算，训练侧闭环仍需进一步实验确认。")
    add_heading(doc, "16.3 面向时空扰动的因果语义回放", 2)
    add_body(doc, "无人机视角、光照、背景、飞行姿态和定位偏差可能诱导模型学习虚假关联。私有方法以因果语义回放为核心，目标是在类别增量学习过程中保持与目标本体相关的稳定语义。公开材料以整体准确率和遗忘率结果验证能力，不披露内部实现。")
    add_heading(doc, "16.4 从诊断到复飞的业务闭环", 2)
    add_body(doc, "项目不仅输出模型指标，还把高风险location_id转换为可执行的复飞坐标、补采清单和人工复核任务。这是北斗时空信息从“记录位置”走向“驱动作业”的关键产品化步骤，也是区别于纯算法论文演示的重要价值。")

    add_heading(doc, "17. 产品价值与可交付物", 1)
    add_table(doc, ["价值对象", "核心价值", "可交付物"], [
        ["巡检操作", "把模型薄弱区域直接转为坐标级复飞和补采任务", "航线任务、复飞清单、异常点位"],
        ["运维管理", "从全局准确率升级为区域风险和历史趋势管理", "风险地图、区域诊断报告、处置状态"],
        ["模型运维", "在有限预算下持续吸收新类别并控制遗忘", "模型版本、质量门槛、回放预算建议"],
        ["管理决策", "量化哪些区域、模型和任务需要优先投入资源", "汇总看板、预算分析、阶段评估"],
    ], [1800, 4000, 3560])
    add_body(doc, "商业化初期不必承诺完全替代人工巡检。更可信的定位是“辅助识别、风险排序和复飞决策平台”，先减少无差别复飞、降低人工筛查量、提高重点区域复核及时性，再随着真实数据和行业适配逐步提高自动化程度。")

    add_heading(doc, "18. 团队协作与保密实现", 1)
    add_picture(doc, isolation, "图4  核心算法与共享竞赛工程通过审核后的白名单JSON连接")
    add_heading(doc, "18.1 推荐分工", 2)
    add_table(doc, ["负责人", "工作内容", "接触的数据"], [
        ["算法负责人", "私有训练、统一评测、结果审核、JSON导出、技术答辩", "源码、配置、权重、预测、日志和正式结果"],
        ["前端/产品成员", "页面、交互、地图、图表、上传流程原型、可访问性", "共享仓库、mock数据和审核后的public_result.json"],
        ["BP/PPT成员", "项目叙事、市场、商业模式、排版、路演与答辩", "说明文档、公开图表、产品截图和可公开指标"],
        ["指导老师/项目负责人", "成果归属、公开边界、参赛与投稿协调", "按需查看完整材料并做最终授权"],
    ], [1900, 3900, 3560])
    add_heading(doc, "18.2 为什么队员的AI仍能理解完整项目", 2)
    add_body(doc, "AI理解项目不需要看到私有算法源码。共享仓库已经包含AGENTS.md、项目上下文、JSON字段契约、页面职责、证据状态和交接文档。队员把beidou_competition_shared作为唯一工作区交给AI，AI就能理解用户、流程、页面、数据字段和允许修改的范围；需要新增指标时，只提出JSON字段需求，由算法负责人决定是否导出。")
    add_heading(doc, "18.3 Git同步方式", 2)
    add_body(doc, "共享仓库用于频繁同步前端、文档和公开结果，不需要每次重新打包。队员提交前端改动，算法负责人拉取；算法负责人跑完实验后只更新data_contract/public_result.json并提交，队员拉取后页面自动读取新结果。私有训练目录、服务器路径、权重和原始预测文件永远不进入共享Git仓库。")

    add_heading(doc, "19. 当前完成状态与缺口", 1)
    add_table(doc, ["能力", "状态", "说明"], [
        ["航线与区域交互仿真", "已完成", "可展示定位扰动、区域风险、预算和复飞闭环，证据为simulation_only"],
        ["公开结果看板与报告", "已完成", "读取审核JSON，保留prediction_backed状态和披露说明"],
        ["正式公开数据实验", "已完成代表性运行", "10类5任务，三层方法优于ER；仍需多随机种子与更充分实验"],
        ["区域风险预算建议", "已完成计算", "固定600总量守恒，当前不等于已控制私有训练缓存"],
        ["真实任务上传与后台队列", "未完成", "当前静态前端尚不能创建真实任务或启动服务器作业"],
        ["真实北斗终端和外业航迹", "未完成", "当前坐标与定位噪声为仿真"],
        ["机载部署与性能测试", "未完成", "没有功耗、延迟、显存、端侧存储和环境适应性指标"],
        ["客户试点与商业验收", "未完成", "不能写成已有客户收入或生产效果"],
    ], [2800, 1500, 5060])

    add_heading(doc, "20. 分阶段路线图", 1)
    add_table(doc, ["阶段", "目标", "主要工作", "验收证据"], [
        ["阶段A：初赛材料", "把方向和证据讲清楚", "完善说明书、BP、产品图、实验表格；保持证据边界", "材料口径一致，正式JSON可追溯"],
        ["阶段B：半决赛原型", "形成可操作演示", "任务上传、字段校验、真实地图、异步作业状态、报告导出", "完整演示流程和录屏"],
        ["阶段C：算法增强", "验证时空风险真正影响训练", "多随机种子、消融、类别平衡、位置风险预算接入私有训练", "均值方差、显著性、区域改善证据"],
        ["阶段D：真实联调", "接入北斗终端和无人机任务包", "真实坐标、时间同步、定位质量、跨架次关联、小规模外业", "真实任务数据与现场复飞闭环"],
        ["阶段E：产品试点", "面向行业客户验证价值", "权限、审计、模型版本、部署、运维和安全", "试点报告、效率指标和客户反馈"],
    ], [1700, 2100, 3560, 2000])

    add_heading(doc, "21. 近期实验应支持哪些结论", 1)
    add_body(doc, "下一轮实验不应只追求一个更高的最终准确率，而应围绕项目主张设计证据链。每个结论都要有明确对照和指标，避免出现结果很多但无法回答评委问题的情况。")
    add_table(doc, ["拟支持结论", "必须补充的实验", "关键指标"], [
        ["三层方法稳定优于ER", "至少3个随机种子、统一任务划分和预算", "平均准确率均值±标准差、遗忘率均值±标准差"],
        ["提升来自三层关键机制", "逐层消融或模块关闭实验", "每个模块带来的增益和遗忘变化"],
        ["位置风险预算有效", "均匀预算 vs. 风险预算，在训练侧真正控制样本/特征曝光", "高风险区域准确率、区域遗忘、总体预算守恒"],
        ["不是只记住最近任务", "任务后准确率矩阵和类别级结果", "旧任务保持率、最近类别偏置、混淆矩阵"],
        ["具备工程可行性", "记录训练/推理时间、显存和缓存占用", "延迟、吞吐、资源占用"],
    ], [3100, 3760, 2500])
    add_callout(doc, "优先级建议", "近期最重要的是多随机种子、任务后准确率矩阵、三层方法与ER的统一对比，以及位置风险预算训练闭环。真实无人机和硬件可以在晋级后逐步补齐。", fill=PALE_GOLD, color=GOLD)

    add_heading(doc, "22. 评委讲解口径", 1)
    add_heading(doc, "22.1 30秒项目介绍", 2)
    add_callout(doc, "标准话术", "我们面向搭载北斗定位模块的无人机常态化巡检场景，解决新类别目标持续出现、模型历史知识遗忘以及不同地理区域识别风险不一致的问题。系统用图像完成目标识别，用北斗坐标建立跨架次时空索引和区域诊断，再依据局部风险生成回放预算、重点复飞和人工复核建议。目前已在公开航拍数据和虚拟北斗航线仿真中完成10类5任务验证，三层方法相较ER准确率提升12.11个百分点、遗忘率下降19.23个百分点。", fill=LIGHT_BLUE)

    add_heading(doc, "22.2 90秒技术介绍", 2)
    add_body(doc, "传统持续学习方法通常只优化全局平均指标，容易忽视高风险地点的局部失效。我们的系统采用双路架构：航拍图像进入视觉识别和私有三层持续学习引擎；北斗坐标、时间和定位质量进入时空索引与区域诊断模块；两路通过样本编号和时间戳关联。系统按location_id统计准确率、遗忘和定位误差，把模型不稳定性与业务后果结合形成区域风险，并在固定预算下给出维护资源和复飞优先级。经纬度暂不直接输入分类器，避免模型学习“某个地点等于某个类别”的位置捷径。当前完成的是公开数据仿真、正式结果看板和诊断报告，真实北斗终端、任务后台和机载部署属于下一阶段。")

    add_heading(doc, "22.3 三条不可越过的表述红线", 2)
    for text in [
        "不要说“已在真实电网或客户现场部署”，除非后续确有外业和验收材料。",
        "不要说“北斗坐标已经直接提升视觉分类精度”，当前坐标主要用于时空诊断、预算建议和复飞定位。",
        "不要把代表性单随机种子结果说成稳定商业指标；正式BP应标明数据集、任务数、预算和仿真边界。",
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "23. 高频疑问与回答", 1)
    faq = [
        ("Q1：北斗只是给图片加经纬度吗？", "不是。经纬度用于跨架次关联、区域划分、局部诊断、风险排序和复飞坐标生成。若只加字段而不影响诊断和作业决策，就只是包装。"),
        ("Q2：为什么不把经纬度直接喂给神经网络？", "当前直接输入容易形成位置捷径并放大定位噪声。双路架构更稳妥：视觉负责识别，北斗负责时空治理，两路在诊断和决策层融合。"),
        ("Q3：这个前端给谁使用？", "完整产品面向巡检操作员、运维管理人员和模型运维人员。当前版本是竞赛原型，重点展示未来操作逻辑和已经审核的实验结果。"),
        ("Q4：用户数据如何进入系统？", "未来由无人机或巡检平台上传标准任务包，包含图像/视频、时间、坐标、定位质量和任务标识；不是让用户逐张输入。"),
        ("Q5：上传后会自动训练吗？", "不会。先推理和诊断，未知类别需要人工确认，训练由模型运维人员受控发起，候选模型通过门槛后再发布。"),
        ("Q6：43.59%是不是太低？", "它是10类5任务持续学习后的最终平均准确率，不能与单任务成熟模型直接比较。当前价值在于同设置下比ER高12.11个百分点且遗忘明显降低；商业部署仍需更多数据和工程验证。"),
        ("Q7：位置风险预算是否已经真正参与训练？", "当前已生成基于区域风险的预算建议和分配结果；是否已反向控制私有训练缓存必须以训练侧联调实验为准，未完成前不得夸大。"),
        ("Q8：队员做前端需要核心算法吗？", "不需要。前端只读取白名单JSON或API。算法负责人保留源码、权重和日志，队员及其AI通过项目上下文和数据契约理解产品。"),
        ("Q9：如果评委要求看代码怎么办？", "可以展示共享工程、数据契约、诊断逻辑和运行结果。私有核心算法处于知识产权和学术成果保护阶段，可说明已完成可重复实验验证，但不现场公开核心源码。"),
        ("Q10：项目最终卖的是什么？", "不是单独出售一个分类模型，而是提供面向巡检任务的数据接入、时空风险诊断、持续模型维护、复飞决策和报告服务，可按软件平台、项目实施或年度运维服务交付。"),
    ]
    for question, answer in faq:
        add_heading(doc, question, 2)
        add_body(doc, answer)

    add_heading(doc, "24. 对外统一项目表述", 1)
    add_callout(doc, "最终版项目陈述", "本项目面向低空无人机常态化巡检，构建北斗时空信息与类别增量持续学习协同的平台。系统将航拍图像与北斗坐标、时间和定位质量通过样本编号关联，以视觉模型识别目标，以时空诊断模块定位不同区域的模型遗忘和识别风险，并在固定资源约束下形成区域回放预算、重点复飞和人工复核建议。私有三层因果语义回放方法用于缓解新类别持续出现带来的灾难性遗忘。当前已完成公开航拍数据、虚拟北斗航线和定位扰动下的仿真验证及竞赛展示原型；后续将补齐风险预算训练闭环、真实任务后台、北斗终端联调和行业试点。", fill=LIGHT_BLUE)

    add_heading(doc, "附录A：公开结果字段解释", 1)
    add_table(doc, ["字段", "解释", "使用注意"], [
        ["evidence_status", "证据状态：simulation_only、position_only或prediction_backed", "不能自动推导外业、部署或商业结果"],
        ["average_accuracy", "连续学习结束后的平均准确率", "必须同时说明任务设置和对比方法"],
        ["average_forgetting", "历史任务性能下降程度", "数值越低通常越好"],
        ["location_id", "地理网格或业务区域标识", "当前区域为仿真构造"],
        ["position_error_mean_m", "区域平均定位误差估计", "当前来自仿真定位扰动"],
        ["risk", "区域模型不稳定性与业务后果形成的综合排序值", "不是事故发生概率"],
        ["replay_budget", "固定总量下的区域预算建议", "不等于训练缓存已自动按此执行"],
        ["recommended_for_bp", "结果是否通过当前自动质量门槛", "仍需人工审核表述和证据边界"],
    ], [2500, 3900, 2960])

    add_heading(doc, "附录B：证据状态与合法表述", 1)
    add_table(doc, ["状态", "含义", "推荐用语"], [
        ["simulation_only", "路线、坐标和指标均为功能仿真", "业务闭环原型、交互仿真、概念验证"],
        ["position_only", "时空元数据链路已处理，但没有模型预测", "时空数据接入和区域索引已跑通"],
        ["prediction_backed", "指标来自模型逐样本预测", "公开数据实验结果；仍需披露仿真与部署边界"],
    ], [1900, 3900, 3560])

    add_heading(doc, "附录C：团队提交前检查清单", 1)
    checks = [
        "项目分类统一写为“北斗时空信息应用—产业与经济应用—低空经济（无人机巡检）”。",
        "所有结果表明确数据集、类别数、任务数、预算和证据状态。",
        "图像识别与北斗时空诊断采用双路架构，经纬度不被误写成当前分类器直接输入。",
        "区分预算建议和训练侧真实预算控制，不把规划写成已实现。",
        "前端截图保留simulation_only或prediction_backed标识及披露说明。",
        "共享仓库不包含私有源码、权重、训练日志、服务器账号和原始预测文件。",
        "商业与社会效益使用预测和规划措辞，不虚构客户、收入、专利或部署数量。",
        "答辩人能够解释43.59%、12.11pp和19.23pp分别代表什么。",
        "所有队员使用同一项目名称、场景、流程和创新点，不各讲一套。",
    ]
    for check in checks:
        add_bullet(doc, check, bullet_id)

    add_callout(doc, "结论", "当前项目方向成立，北斗关系可以讲清，公开数据结果已具备初赛材料候选价值。下一阶段的关键不是继续堆砌页面，而是用多随机种子、消融和训练侧位置风险预算闭环增强证据，同时逐步把静态展示页升级为可操作的巡检任务原型。", fill=PALE_GOLD, color=GOLD)

    # Core properties and deterministic save.
    props = doc.core_properties
    props.title = "北斗巡知项目完整说明书"
    props.subject = "北斗时空信息与低空无人机巡检项目团队统一说明"
    props.author = "项目组"
    props.keywords = "北斗, 无人机巡检, 类别增量持续学习, 时空风险诊断"
    props.comments = "基于公开竞赛工程和审核后的公开结果生成。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
